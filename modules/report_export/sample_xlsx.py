"""Controlled Excel representation of the effective market sample."""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from collections.abc import Mapping
from datetime import datetime, timezone
from numbers import Real
from typing import Any, Dict, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from ..provenance import canonical_json
from ..results_generator import build_report_view

_FIXED_TIME = datetime(2020, 1, 1, tzinfo=timezone.utc)
_FIXED_ZIP_TIME = (2020, 1, 1, 0, 0, 0)
_FIXED_W3CDTF = b"2020-01-01T00:00:00Z"


def _cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (bool, int, float)):
        return value
    text = str(value)
    if text == "":
        return None
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _deterministic_zip(data: bytes) -> bytes:
    """Normalize OOXML member order and timestamps for stable byte hashes."""
    source = zipfile.ZipFile(io.BytesIO(data), "r")
    output = io.BytesIO()
    with source, zipfile.ZipFile(
        output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as target:
        for name in sorted(source.namelist()):
            original = source.getinfo(name)
            info = zipfile.ZipInfo(name, date_time=_FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = original.external_attr
            info.create_system = original.create_system
            payload = source.read(name)
            if name == "docProps/core.xml":
                # openpyxl overwrites workbook.properties.modified with the
                # wall clock at save time.  ZIP metadata normalization alone
                # therefore cannot make repeated builds byte-identical.
                payload, replacements = re.subn(
                    rb"(<dcterms:modified\b[^>]*>)[^<]*(</dcterms:modified>)",
                    rb"\g<1>" + _FIXED_W3CDTF + rb"\g<2>",
                    payload,
                    count=1,
                )
                if replacements != 1:
                    raise ValueError("OOXML core properties lack one modified timestamp")
            target.writestr(info, payload)
    return output.getvalue()


def build_sample_xlsx(
    snapshot: Mapping[str, Any],
    report_context: Optional[Mapping[str, Any]] = None,
) -> bytes:
    """Emit the full effective sample with source and decimal coordinates."""
    view = build_report_view(snapshot, report_context)
    workbook = Workbook()
    workbook.properties.creator = "MODELA PRO"
    workbook.properties.created = _FIXED_TIME
    workbook.properties.modified = _FIXED_TIME
    sample = workbook.active
    sample.title = "amostra_efetiva"
    headers = [
        "nº", "row_id", "endereço_completo", "latitude_decimal",
        "longitude_decimal", "fonte", "justificativa",
    ] + list(view.get("used_value_columns") or [])
    sample.append([_cell(value) for value in headers])
    for cell in sample[1]:
        cell.font = Font(bold=True)
    for row in view.get("used_rows") or []:
        sample.append([_cell(value) for value in [
            row.get("seq"), row.get("row_id"), row.get("address"),
            row.get("latitude"), row.get("longitude"), row.get("source"),
            row.get("justification"),
        ] + list(row.get("value_cells") or [])])
    sample.freeze_panes = "A2"
    sample.auto_filter.ref = sample.dimensions

    excluded = workbook.create_sheet("dados_excluidos")
    excluded_headers = [
        "nº", "row_id", "endereço_completo", "latitude_decimal",
        "longitude_decimal", "fonte", "justificativa",
    ] + list(view.get("excluded_value_columns") or [])
    excluded.append([_cell(value) for value in excluded_headers])
    for cell in excluded[1]:
        cell.font = Font(bold=True)
    for row in view.get("excluded_rows") or []:
        excluded.append([_cell(value) for value in [
            row.get("seq"), row.get("row_id"), row.get("address"),
            row.get("latitude"), row.get("longitude"), row.get("source"),
            row.get("justification"),
        ] + list(row.get("value_cells") or [])])
    excluded.freeze_panes = "A2"
    excluded.auto_filter.ref = excluded.dimensions

    subject = workbook.create_sheet("imovel_avaliando")
    subject.append([_cell("campo"), _cell("valor")])
    subject["A1"].font = subject["B1"].font = Font(bold=True)
    for key, value in view.get("subject_items") or []:
        if isinstance(value, Mapping):
            continue
        subject.append([_cell(key), _cell(value)])
    location = view.get("subject_geolocation") or {}
    for key in ("address", "latitude", "longitude", "source", "coordinate_system"):
        subject.append([_cell("geolocalizacao." + key), _cell(location.get(key))])

    variables = workbook.create_sheet("criterios_variaveis")
    variables.append([
        _cell(value) for value in
        ["variavel", "criterio_enquadramento", "codificacao", "categorias", "escala"]
    ])
    for cell in variables[1]:
        cell.font = Font(bold=True)
    for row in view.get("variable_classification") or []:
        variables.append([
            _cell(row.get("variable")), _cell(row.get("criterion")),
            _cell(row.get("coding")), _cell(row.get("categories")),
            _cell(row.get("scale_values")),
        ])

    control = workbook.create_sheet("controle")
    row_ids = [str(row.get("row_id")) for row in view.get("used_rows") or []]
    control.append([_cell("schema_version"), _cell("MP-SAMPLE-XLSX/1")])
    control.append([_cell("job_id"), _cell(str(snapshot.get("job_id") or ""))])
    control.append([_cell("effective_row_count"), _cell(len(row_ids))])
    control.append([
        _cell("effective_row_ids_sha256"),
        _cell(hashlib.sha256(canonical_json(row_ids).encode("utf-8")).hexdigest()),
    ])
    output = io.BytesIO()
    workbook.save(output)
    return _deterministic_zip(output.getvalue())


def _typed_cells_equal(observed: Any, expected: Any) -> bool:
    if isinstance(observed, bool) or isinstance(expected, bool):
        return isinstance(observed, bool) and isinstance(expected, bool) and observed == expected
    if isinstance(observed, Real) and isinstance(expected, Real):
        return float(observed) == float(expected)
    return type(observed) is type(expected) and observed == expected


def _typed_rows_equal(observed: Any, expected: Any) -> bool:
    if len(observed) != len(expected):
        return False
    return all(
        len(observed_row) == len(expected_row)
        and all(
            _typed_cells_equal(observed_cell, expected_cell)
            for observed_cell, expected_cell in zip(observed_row, expected_row)
        )
        for observed_row, expected_row in zip(observed, expected)
    )


def verify_sample_xlsx(
    data: bytes,
    snapshot: Mapping[str, Any],
    report_context: Optional[Mapping[str, Any]] = None,
) -> Dict[str, Any]:
    """Compare the workbook's exact effective row IDs and required fields."""
    findings = []
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=False, data_only=False)
    except Exception as exc:
        return {"ok": False, "findings": [{"code": "SAMPLE_XLSX_INVALID", "message": str(exc)}]}
    required = {
        "amostra_efetiva", "dados_excluidos", "imovel_avaliando",
        "criterios_variaveis", "controle",
    }
    if not required.issubset(workbook.sheetnames):
        findings.append({"code": "SAMPLE_XLSX_SHEET_MISSING", "missing": sorted(required - set(workbook.sheetnames))})
        return {"ok": False, "findings": findings}
    for sheet_name in sorted(required):
        required_sheet = workbook[sheet_name]
        if required_sheet.sheet_state != "visible":
            findings.append({"code": "SAMPLE_XLSX_SHEET_HIDDEN", "sheet": sheet_name})
        for index in range(1, required_sheet.max_row + 1):
            if required_sheet.row_dimensions[index].hidden is True:
                findings.append({
                    "code": "SAMPLE_XLSX_ROW_HIDDEN", "sheet": sheet_name,
                    "row": index,
                })
        for index in range(1, required_sheet.max_column + 1):
            letter = required_sheet.cell(1, index).column_letter
            if required_sheet.column_dimensions[letter].hidden is True:
                findings.append({
                    "code": "SAMPLE_XLSX_COLUMN_HIDDEN", "sheet": sheet_name,
                    "column": letter,
                })
        for row in required_sheet.iter_rows():
            for cell in row:
                if cell.number_format != "General":
                    findings.append({
                        "code": "SAMPLE_XLSX_CELL_FORMAT_UNEXPECTED",
                        "sheet": sheet_name, "cell": cell.coordinate,
                        "number_format": cell.number_format,
                    })
    view = build_report_view(snapshot, report_context)
    formula_cells = [
        f"{sheet.title}!{cell.coordinate}"
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.data_type == "f"
    ]
    if formula_cells:
        findings.append({"code": "SAMPLE_XLSX_FORMULA_CELL", "cells": formula_cells})
    sheet = workbook["amostra_efetiva"]
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    headers = list(rows[0]) if rows else []
    expected_headers = [
        "nº", "row_id", "endereço_completo", "latitude_decimal",
        "longitude_decimal", "fonte", "justificativa",
    ] + list(view.get("used_value_columns") or [])
    expected_headers = [_cell(value) for value in expected_headers]
    if headers != expected_headers:
        findings.append({"code": "SAMPLE_XLSX_HEADERS_MISMATCH"})
    expected_ids = [str(row.get("row_id")) for row in view.get("used_rows") or []]
    observed_ids = [str(row[1]) for row in rows[1:]]
    if observed_ids != expected_ids:
        findings.append({"code": "SAMPLE_XLSX_ROW_IDS_MISMATCH"})
    expected_rows = [expected_headers]
    for row in view.get("used_rows") or []:
        expected_rows.append([_cell(value) for value in [
            row.get("seq"), row.get("row_id"), row.get("address"),
            row.get("latitude"), row.get("longitude"), row.get("source"),
            row.get("justification"),
        ] + list(row.get("value_cells") or [])])
    if not _typed_rows_equal(rows, expected_rows):
        findings.append({"code": "SAMPLE_XLSX_CONTENT_MISMATCH"})

    excluded_rows = [
        list(row) for row in workbook["dados_excluidos"].iter_rows(values_only=True)
    ]
    expected_excluded_headers = [_cell(value) for value in [
        "nº", "row_id", "endereço_completo", "latitude_decimal",
        "longitude_decimal", "fonte", "justificativa",
    ] + list(view.get("excluded_value_columns") or [])]
    expected_excluded_rows = [expected_excluded_headers]
    for row in view.get("excluded_rows") or []:
        expected_excluded_rows.append([_cell(value) for value in [
            row.get("seq"), row.get("row_id"), row.get("address"),
            row.get("latitude"), row.get("longitude"), row.get("source"),
            row.get("justification"),
        ] + list(row.get("value_cells") or [])])
    if not _typed_rows_equal(excluded_rows, expected_excluded_rows):
        findings.append({"code": "SAMPLE_XLSX_EXCLUDED_CONTENT_MISMATCH"})

    subject_rows = [
        list(row) for row in workbook["imovel_avaliando"].iter_rows(values_only=True)
    ]
    expected_subject = [[_cell("campo"), _cell("valor")]]
    for key, value in view.get("subject_items") or []:
        if not isinstance(value, Mapping):
            expected_subject.append([_cell(key), _cell(value)])
    location = view.get("subject_geolocation") or {}
    for key in ("address", "latitude", "longitude", "source", "coordinate_system"):
        expected_subject.append([_cell("geolocalizacao." + key), _cell(location.get(key))])
    if not _typed_rows_equal(subject_rows, expected_subject):
        findings.append({"code": "SAMPLE_XLSX_SUBJECT_MISMATCH"})

    criteria_rows = [
        list(row) for row in workbook["criterios_variaveis"].iter_rows(values_only=True)
    ]
    expected_criteria = [[_cell(value) for value in [
        "variavel", "criterio_enquadramento", "codificacao", "categorias", "escala"
    ]]]
    for row in view.get("variable_classification") or []:
        expected_criteria.append([_cell(row.get(key)) for key in (
            "variable", "criterion", "coding", "categories", "scale_values"
        )])
    if not _typed_rows_equal(criteria_rows, expected_criteria):
        findings.append({"code": "SAMPLE_XLSX_CRITERIA_MISMATCH"})

    expected_control = [
        [_cell("schema_version"), _cell("MP-SAMPLE-XLSX/1")],
        [_cell("job_id"), _cell(str(snapshot.get("job_id") or ""))],
        [_cell("effective_row_count"), _cell(len(expected_ids))],
        [_cell("effective_row_ids_sha256"), _cell(
            hashlib.sha256(canonical_json(expected_ids).encode("utf-8")).hexdigest()
        )],
    ]
    control_rows = [
        list(row) for row in workbook["controle"].iter_rows(values_only=True)
    ]
    if not _typed_rows_equal(control_rows, expected_control):
        findings.append({"code": "SAMPLE_XLSX_CONTROL_MISMATCH"})
    all_document_rows = rows[1:] + excluded_rows[1:]
    geolocation_complete = bool(rows[1:]) and not any(
        len(row) < 7 or row[2] in (None, "") or row[3] is None
        or row[4] is None or row[5] in (None, "")
        or not isinstance(row[3], (int, float)) or isinstance(row[3], bool)
        or not isinstance(row[4], (int, float)) or isinstance(row[4], bool)
        or not -90 <= float(row[3]) <= 90
        or not -180 <= float(row[4]) <= 180
        for row in all_document_rows
    )
    expected_geolocation_complete = bool(expected_ids) and all(
        row.get("geolocation_complete") is True
        for row in list(view.get("used_rows") or [])
        + list(view.get("excluded_rows") or [])
    )
    if expected_geolocation_complete and not geolocation_complete:
        findings.append({"code": "SAMPLE_XLSX_GEOLOCATION_INVALID"})
    return {
        "ok": not findings,
        "findings": findings,
        "sha256": hashlib.sha256(data).hexdigest(),
        "size": len(data),
        "row_count": len(observed_ids),
        "geolocation_complete": geolocation_complete,
    }
