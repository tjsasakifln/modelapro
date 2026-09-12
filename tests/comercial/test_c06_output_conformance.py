"""Production representation contract for C06 documentary conformance.

All people, properties, market rows and certificate claims are synthetic test
data. They prove product capability only; they are not an appraisal or an
ICP-Brasil act.
"""

from __future__ import annotations

import copy
import hashlib
import io
import json
import re
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook

from backend.worker import compose_valuation_job, resolve_peers
from modules.job_store import JobStore
from modules.digital_signatures.pdf import (
    ICP_BRASIL_PADES_POLICY_REGISTRY,
    ICP_BRASIL_TRUST_ANCHORS,
    _signature_policy_details,
    _signature_policy_verified,
)
from modules.digital_signatures import (
    prepare_signature_request,
    record_external_signature,
    verify_pdf_signature,
    verify_signature_binding,
)
from modules.pro_workflow.report_context import (
    build_output_manifest,
    complete_report_context,
)
from modules.qualification_profile import resolve_profile
from modules.qualification_profile.output_conformance import (
    assess_output_conformance,
    product_conformance_baseline,
)
from modules.report_export import build_docx, build_sample_xlsx, verify_sample_xlsx
from modules.report_export.workflow import (
    _representation_manifest,
    _verify_representation_manifest,
)
from modules.report_presenter.qualification import (
    _normalize_external_signature_state,
    report_content_fingerprint,
    signable_snapshot_sha256,
)
from modules.report_presenter.verifier import extract_pdf_text, verify_report_consistency
from modules.results_generator import build_report_view, render_report
from tests.comercial.c03.fixtures import qualified_case
from tests.comercial.test_c06_numeric_disclosure import _noisy_linear_csv
from tests.comercial.test_c06_document_flow import _professional_spec


def _signed_with_crl(pdf: bytes, *, revoked: bool = False):
    """Create a generic TESTE chain and a real signed CRL; never an ICP claim."""
    from asn1crypto import crl as asn1_crl
    from asn1crypto import x509 as asn1_x509
    from cryptography import x509
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import rsa
    from cryptography.hazmat.primitives.serialization import pkcs12
    from cryptography.x509.oid import NameOID
    from pyhanko.pdf_utils.incremental_writer import IncrementalPdfFileWriter
    from pyhanko.sign import signers
    from pyhanko_certvalidator import ValidationContext

    now = datetime.now(timezone.utc)
    ca_key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    ca_name = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, "CA SINTETICA TESTE")])
    ca = (
        x509.CertificateBuilder().subject_name(ca_name).issuer_name(ca_name)
        .public_key(ca_key.public_key()).serial_number(1)
        .not_valid_before(now - timedelta(days=2)).not_valid_after(now + timedelta(days=10))
        .add_extension(x509.BasicConstraints(ca=True, path_length=None), critical=True)
        .add_extension(x509.KeyUsage(
            digital_signature=True, content_commitment=False, key_encipherment=False,
            data_encipherment=False, key_agreement=False, key_cert_sign=True,
            crl_sign=True, encipher_only=False, decipher_only=False,
        ), critical=True).sign(ca_key, hashes.SHA256())
    )
    signer_key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    signer_name = x509.Name([
        x509.NameAttribute(NameOID.COMMON_NAME, "SIGNATARIO SINTETICO TESTE")
    ])
    signer_cert = (
        x509.CertificateBuilder().subject_name(signer_name).issuer_name(ca_name)
        .public_key(signer_key.public_key()).serial_number(2)
        .not_valid_before(now - timedelta(days=1)).not_valid_after(now + timedelta(days=5))
        .add_extension(x509.BasicConstraints(ca=False, path_length=None), critical=True)
        .add_extension(x509.KeyUsage(
            digital_signature=True, content_commitment=True, key_encipherment=False,
            data_encipherment=False, key_agreement=False, key_cert_sign=False,
            crl_sign=False, encipher_only=False, decipher_only=False,
        ), critical=True).sign(ca_key, hashes.SHA256())
    )
    builder = (
        x509.CertificateRevocationListBuilder().issuer_name(ca_name)
        .last_update(now - timedelta(minutes=1)).next_update(now + timedelta(days=1))
    )
    if revoked:
        builder = builder.add_revoked_certificate(
            x509.RevokedCertificateBuilder().serial_number(2)
            .revocation_date(now - timedelta(minutes=2)).build()
        )
    crl = builder.sign(ca_key, hashes.SHA256())
    p12 = pkcs12.serialize_key_and_certificates(
        b"TESTE", signer_key, signer_cert, [ca],
        serialization.BestAvailableEncryption(b"senha-teste"),
    )
    signer = signers.SimpleSigner.load_pkcs12_data(
        p12, other_certs=[], passphrase=b"senha-teste"
    )
    signed = signers.PdfSigner(
        signers.PdfSignatureMetadata(field_name="AssinaturaTeste"), signer=signer
    ).sign_pdf(IncrementalPdfFileWriter(io.BytesIO(pdf))).getvalue()
    context = ValidationContext(
        trust_roots=[asn1_x509.Certificate.load(
            ca.public_bytes(serialization.Encoding.DER)
        )],
        crls=[asn1_crl.CertificateList.load(
            crl.public_bytes(serialization.Encoding.DER)
        )],
        revocation_mode="require", allow_fetching=False,
    )
    return signed, context


def _rich_case():
    snapshot, base = qualified_case()
    snapshot = copy.deepcopy(snapshot)
    base = copy.deepcopy(base)
    used_ids = snapshot["sample"]["used_row_ids"][:3]
    excluded_ids = snapshot["sample"]["excluded_row_ids"][:1]
    snapshot["sample"].update({
        "received": 4, "observed_target": 4, "prepared": 4,
        "used": 3, "used_row_ids": used_ids,
        "excluded": 1, "excluded_row_ids": excluded_ids,
    })
    base["used_rows"] = base["used_rows"][:3]
    base["excluded_rows"] = base["excluded_rows"][:1]
    base["fitted_values"] = [12.0, 12.1, 12.2]
    base["residuals"] = [-0.1, 0.0, 0.1]
    base["observed_values"] = [11.9, 12.1, 12.3]
    base["series_row_ids"] = used_ids
    base["series_scale"] = "log"
    snapshot["validation"]["fundamentacao"].update({
        "grade": 3,
        "items": [{
            "item": 1, "description": "Caracterização sintética",
            "grade": 3, "points": 3, "detail": "TESTE",
            "evidence_status": "declared",
        }],
        "points": 18,
    })
    snapshot["model"].pop("formula", None)
    snapshot["model"]["metrics"].update({"r": 0.9539392014})
    snapshot["model"]["target_transform_state"] = {
        "name": "log",
        "inverse": {
            "default_method": "exponential",
            "default_estimand": "mediana condicional na unidade original",
        },
    }
    snapshot["validation"]["statistical"]["diagnostics"] = {
        "standardized_residuals": {
            "available": True, "reason": None,
            "definition": "raw_residual/sample_residual_std_ddof_1",
            "values": [-1.0, 0.0, 1.0], "n": 3,
        },
        "normal_frequency_comparison": {
            "available": True, "reason": None, "n": 3,
            "intervals": [
                {"z": 1.0, "nominal_probability": .68, "nominal_percent": 68,
                 "observed_count": 3, "observed_probability": 1.0},
                {"z": 1.64, "nominal_probability": .90, "nominal_percent": 90,
                 "observed_count": 3, "observed_probability": 1.0},
                {"z": 1.96, "nominal_probability": .95, "nominal_percent": 95,
                 "observed_count": 3, "observed_probability": 1.0},
            ],
        },
        "correlation_matrix": {
            "available": True, "reason": None, "sample": "effective_model_sample",
            "status": "complete", "missing_cells": 0, "missing_variables": [],
            "scale": "original", "variables": ["preco", "area"],
            "values": [[1.0, .9], [.9, 1.0]], "n": 3,
            "design": {"sample": "effective_model_sample", "scale": "design",
                       "variables": ["area"], "values": [[1.0]], "n": 3},
        },
        "elasticities": {
            "available": True, "reason": None,
            "method": "central_finite_difference_predict_original",
            "point": "subject", "relative_step": 1e-4,
            "items": [{"variable": "area", "elasticity": .52,
                       "base_value": 120, "base_prediction": 350000,
                       "step": .012, "derivative": 1516.67}],
            "warnings": [],
            "coverage": {"expected_quantitative_variables": ["area"],
                         "computed_variables": ["area"], "missing_variables": [],
                         "complete": True},
        },
        "outlier_count": {
            "available": True, "reason": None, "detected": 1,
            "excluded": 1, "influential": 0,
            "definition": {"outlier": "|standardized_residual| > 2"},
            "coverage": {"complete": True},
        },
    }
    evidence = {
        row_id: {
            "address": f"Rua de Teste, {index}, Cidade Sintética/TS",
            "latitude": -27.5 - index / 1000,
            "longitude": -48.5 - index / 1000,
            "source": f"TESTE: anúncio sintético {index}",
        }
        for index, row_id in enumerate(used_ids, start=1)
    }
    evidence[excluded_ids[0]] = {
        "address": "Rua Excluída, 1, Cidade Sintética/TS",
        "latitude": -27.61,
        "longitude": -48.61,
        "source": "TESTE: registro excluído preservado",
    }
    report_fields = {
        "objective": "Estimar o valor de mercado do imóvel sintético.",
        "market_diagnosis": "Mercado sintético com oferta estável para este teste.",
        "variable_classification": {
            "area": {"criterion": "área privativa em m²", "coding": "quantitativa contínua"},
            "quartos": {"criterion": "quantidade observada", "scale_values": [1, 2, 3]},
        },
        "observations": "Intervalo admissível registrado no campo de observações.",
        "grade_i_justification": "Não aplicável: o caso sintético atingiu Grau III.",
        "professional_identity": {
            "name": "Profissional Sintético de Teste", "council": "CREA-TESTE",
            "registration": "000000", "art_rrt": "ART-TESTE-000",
        },
        "subject": {
            **base["subject"],
            "geolocation": {
                "address": "Rua do Avaliando, 1, Cidade Sintética/TS",
                "latitude": -27.59, "longitude": -48.55,
                "source": "TESTE: vistoria sintética declarada",
            },
        },
        "sample_evidence": evidence,
    }
    snapshot["value"]["adopted_value"]["reason"] = (
        "Ponto central adotado na fixture sintética."
    )
    context = complete_report_context(
        base,
        request_spec={"report_context": report_fields},
        subject_raw=report_fields["subject"],
        snapshot=snapshot,
    )
    qctx = snapshot["provenance"]["qualification_context"]
    qctx["review_events"][0]["report_content_fingerprint"] = (
        report_content_fingerprint(snapshot, context)
    )
    return snapshot, context


def test_real_pdf_docx_xlsx_and_manifest_cover_every_applicable_bb_item_except_signature():
    snapshot, context = _rich_case()
    xlsx = build_sample_xlsx(snapshot, context)
    xlsx_check = verify_sample_xlsx(xlsx, snapshot, context)
    assert xlsx_check["ok"] is True
    assert xlsx_check["row_count"] == 3
    assert xlsx_check["geolocation_complete"] is True
    assert xlsx == build_sample_xlsx(snapshot, context)
    workbook = load_workbook(io.BytesIO(xlsx), read_only=True)
    rows = list(workbook["amostra_efetiva"].iter_rows(values_only=True))
    assert rows[1][1:6] == (
        "used-0001", "Rua de Teste, 1, Cidade Sintética/TS",
        -27.501, -48.501, "TESTE: anúncio sintético 1",
    )
    excluded_rows = list(
        workbook["dados_excluidos"].iter_rows(values_only=True)
    )
    assert excluded_rows[1][1:6] == (
        snapshot["sample"]["excluded_row_ids"][0],
        "Rua Excluída, 1, Cidade Sintética/TS", -27.61, -48.61,
        "TESTE: registro excluído preservado",
    )

    xlsx_meta = {
        "sha256": hashlib.sha256(xlsx).hexdigest(), "size": len(xlsx),
        "verified": True,
    }
    context["output_manifest"] = build_output_manifest(
        snapshot, context, representations={"sample.xlsx": xlsx_meta}
    )
    pdf = render_report(snapshot, context)
    docx = build_docx(snapshot, context)
    representation = _representation_manifest(
        snapshot, context, pdf=pdf, docx=docx, sample_xlsx=xlsx
    )
    assert _verify_representation_manifest(
        representation, snapshot, context,
        {"report.pdf": pdf, "report.docx": docx, "sample.xlsx": xlsx},
    )["ok"] is True

    pdf_text = extract_pdf_text(pdf)
    assert "Objetivo da avaliação" in pdf_text
    assert "Diagnóstico do mercado" in pdf_text
    assert "variável dependente na unidade original" in pdf_text
    assert "central_finite_difference_predict_original" in pdf_text
    with zipfile.ZipFile(io.BytesIO(docx)) as archive:
        docx_xml = archive.read("word/document.xml").decode("utf-8")
    assert "Objetivo da avaliação" in docx_xml
    assert "Equação na unidade original" in docx_xml
    assert "Matriz de correlações" in docx_xml

    profile = resolve_profile({"id": "bb-meci-avaliacao-imovel-pf", "version": "0.3.0"})
    assessment = assess_output_conformance(profile, representation)
    assert assessment["product_gaps"] == []
    assert assessment["awaiting_human"] == []
    assert assessment["awaiting_signature"] == ["bb.guiar.assinatura_icp"]
    assert assessment["not_applicable"] == ["meci.3.4.1.justificativa_grau_i"]
    assert {item["requirement_id"] for item in assessment["blocking"]} == {
        "bb.guiar.assinatura_icp"
    }


def test_manifest_rejects_labels_missing_diagnostics_and_tampered_bytes():
    snapshot, context = _rich_case()
    context["output_evidence"] = {
        "meci.3.3.1.k": True,
        "bb.guiar.assinatura_icp": "cliente disse que é ICP",
    }
    snapshot["validation"]["statistical"]["diagnostics"][
        "normal_frequency_comparison"
    ] = {"available": False, "reason": "insufficient_sample", "intervals": []}
    xlsx = build_sample_xlsx(snapshot, context)
    manifest = build_output_manifest(snapshot, context, representations={
        "sample.xlsx": {"verified": True, "sha256": hashlib.sha256(xlsx).hexdigest(),
                        "size": len(xlsx)}
    })
    assert "meci.3.3.1.k" not in manifest["items"]
    assert "bb.guiar.assinatura_icp" not in manifest["items"]
    assert manifest["unverified_declarations"] == context["output_evidence"]

    context["output_manifest"] = manifest
    pdf = render_report(snapshot, context)
    docx = build_docx(snapshot, context)
    representation = _representation_manifest(
        snapshot, context, pdf=pdf, docx=docx, sample_xlsx=xlsx
    )
    check = _verify_representation_manifest(
        representation, snapshot, context,
        {"report.pdf": pdf + b"tamper", "report.docx": docx, "sample.xlsx": xlsx},
    )
    assert check["ok"] is False
    assert any(item["code"] == "OUTPUT_REPRESENTATION_BYTES_MISMATCH"
               for item in check["findings"])
    altered_manifest = copy.deepcopy(representation)
    altered_manifest["items"]["meci.3.3.1.k"] = "declaração adulterada"
    check = _verify_representation_manifest(
        altered_manifest, snapshot, context,
        {"report.pdf": pdf, "report.docx": docx, "sample.xlsx": xlsx},
    )
    assert any(item["code"] == "OUTPUT_MANIFEST_DERIVATION_MISMATCH"
               for item in check["findings"])


def test_sample_row_override_wins_mapping_and_persists_for_reopening():
    base = {
        "used_rows": [{
            "row_id": "R000000",
            "values": {
                "id": "CSV-17",
                "endereco": "Rua Planilha, 17",
                "lat": -27.51,
                "lon": -48.51,
                "fonte": "TESTE: planilha original",
            },
            "geolocation": {
                "address": "Rua Anterior, 1",
                "latitude": -27.1,
                "longitude": -48.1,
                "source": "TESTE: evidência anterior",
            },
        }],
        "excluded_rows": [],
    }
    context = complete_report_context(
        base,
        request_spec={
            "report_context": {
                "sample_evidence_columns": {
                    "address": "endereco", "latitude": "lat",
                    "longitude": "lon", "source": "fonte",
                },
                "sample_evidence": {
                    "R000000": {
                        "address": "Rua Vistoria, 99",
                        "latitude": -27.99,
                        "longitude": -48.99,
                        "source": "TESTE: vistoria profissional",
                    }
                },
            }
        },
        snapshot={"input_sha256": "a" * 64},
    )
    row = context["used_rows"][0]
    assert row["row_id"] == "R000000"
    assert row["values"]["id"] == "CSV-17"
    assert row["values"]["fonte"] == "TESTE: planilha original"
    assert row["geolocation"] == {
        "latitude": -27.99,
        "longitude": -48.99,
        "address": "Rua Vistoria, 99",
        "source": "TESTE: vistoria profissional",
        "coordinate_system": "decimal_degrees_wgs84",
        "complete": True,
        "issues": [],
    }
    assert context["sample_evidence"]["R000000"] == {
        "address": "Rua Vistoria, 99",
        "latitude": -27.99,
        "longitude": -48.99,
        "source": "TESTE: vistoria profissional",
    }


def test_wide_correlation_is_split_without_losing_values_and_pvalues_stay_nonzero():
    snapshot, context = _rich_case()
    count = 20
    variables = ["preco"] + [f"variavel_{index:02d}" for index in range(1, count)]
    values = [
        [1.0 if row == column else 0.8765 for column in range(count)]
        for row in range(count)
    ]
    matrix = snapshot["validation"]["statistical"]["diagnostics"][
        "correlation_matrix"
    ]
    matrix.update(variables=variables, values=values)
    snapshot["model"]["metrics"]["f_pvalue"] = 1e-12
    snapshot["model"]["pvalues"] = {"const": 0.05, "area": 1e-12}
    view = build_report_view(snapshot, context)
    rendered = view["statistical_diagnostics"]["correlation"]
    assert [len(panel["variables"]) for panel in rendered["panels"]] == [6, 6, 6, 2]
    reconstructed = [[] for _ in range(count)]
    for panel in rendered["panels"]:
        assert len(panel["rows"]) == count
        for index, row in enumerate(panel["rows"]):
            assert row["variable"] == variables[index]
            reconstructed[index].extend(row["values"])
    expected = [[f"{value:.4f}" for value in row] for row in values]
    assert reconstructed == expected
    assert view["metrics"]["f_pvalue"] == "1.000e-12"
    pvalues = {row["variable"]: row["pvalue"] for row in view["coef_rows"]}
    assert pvalues["const"] == "0.0500"
    assert pvalues["area"] == "1.000e-12"

    sample_columns = [f"campo_{index}" for index in range(12)]
    sample_rows = [{
        "seq": index + 1,
        "row_id": f"R{index:06d}",
        "value_cells": [str(index * 100 + column) for column in range(12)],
    } for index in range(3)]
    from modules.results_generator import _ledger_value_panels
    sample_panels = _ledger_value_panels(sample_rows, sample_columns, kind="used")
    assert [len(panel["columns"]) for panel in sample_panels] == [5, 5, 2]
    for index, row in enumerate(sample_rows):
        reconstructed_row = []
        for panel in sample_panels:
            assert panel["rows"][index]["row_id"] == row["row_id"]
            reconstructed_row.extend(panel["rows"][index]["cells"])
        assert reconstructed_row == row["value_cells"]

    pdf_text = extract_pdf_text(render_report(snapshot, context))
    assert "Painel 4 de 4" in pdf_text
    assert "0.8765" in pdf_text
    with zipfile.ZipFile(io.BytesIO(build_docx(snapshot, context))) as archive:
        docx_xml = archive.read("word/document.xml").decode("utf-8")
    assert "Painel 4 de 4" in docx_xml
    assert "0.8765" in docx_xml


def test_sample_panel_verifier_binds_each_value_to_row_and_panel():
    snapshot, context = _rich_case()
    for row in context["used_rows"]:
        row["values"]["campo_extra"] = row["label"]
    pdf = render_report(snapshot, context)
    text = extract_pdf_text(pdf)
    view = build_report_view(snapshot, context)
    assert len(view["used_value_panels"]) == 2
    assert verify_report_consistency(
        pdf, snapshot, context, extracted_text=text
    )["ok"] is True

    first = view["used_value_panels"][0]["rows"][0]
    second = view["used_value_panels"][0]["rows"][1]
    first_pair = f"{first['cells'][0]}\n{first['cell_tokens'][0]}"
    second_pair = f"{second['cells'][0]}\n{second['cell_tokens'][0]}"
    assert first_pair in text and second_pair in text

    swapped = text.replace(
        first_pair, f"{second['cells'][0]}\n{first['cell_tokens'][0]}", 1
    ).replace(
        second_pair, f"{first['cells'][0]}\n{second['cell_tokens'][0]}", 1
    )
    omitted = text.replace(first_pair, "", 1)
    duplicated = text.replace(first_pair, first_pair + "\n" + first_pair, 1)
    other_panel = view["used_value_panels"][1]["rows"][0]
    same_value_wrong_panel = text.replace(
        first_pair, f"{first['cells'][0]}\n{other_panel['cell_tokens'][0]}", 1
    )
    for mutation in (swapped, omitted, duplicated, same_value_wrong_panel):
        check = verify_report_consistency(
            pdf, snapshot, context, extracted_text=mutation
        )
        assert check["ok"] is False
        assert any(item["code"] == "MUTATED_SAMPLE_CELL"
                   for item in check["findings"])
    panel = view["used_value_panels"][0]
    first_header = f"{panel['columns'][0]}\n{panel['header_tokens'][0]}"
    second_header = f"{panel['columns'][1]}\n{panel['header_tokens'][1]}"
    swapped_headers = text.replace(
        first_header, f"{panel['columns'][1]}\n{panel['header_tokens'][0]}", 1
    ).replace(
        second_header, f"{panel['columns'][0]}\n{panel['header_tokens'][1]}", 1
    )
    first_row_id = f"{first['row_id']}\n{first['row_id_token']}"
    second_row_id = f"{second['row_id']}\n{second['row_id_token']}"
    swapped_row_ids = text.replace(
        first_row_id, f"{second['row_id']}\n{first['row_id_token']}", 1
    ).replace(
        second_row_id, f"{first['row_id']}\n{second['row_id_token']}", 1
    )
    for mutation, code in (
        (swapped_headers, "MUTATED_SAMPLE_PANEL_HEADER"),
        (swapped_row_ids, "MUTATED_SAMPLE_PANEL_ROW_ID"),
    ):
        check = verify_report_consistency(
            pdf, snapshot, context, extracted_text=mutation
        )
        assert check["ok"] is False
        assert any(item["code"] == code for item in check["findings"])


def test_pdf_verifier_binds_metrics_and_diagnostics_to_their_sections():
    snapshot, context = _rich_case()
    pdf = render_report(snapshot, context)
    text = extract_pdf_text(pdf)
    view = build_report_view(snapshot, context)
    assert verify_report_consistency(
        pdf, snapshot, context, extracted_text=text
    )["ok"] is True
    metrics = view["metrics"]
    diagnostics = view["statistical_diagnostics"]
    normal = diagnostics["normal_frequency"]["rows"][0]
    matrix = diagnostics["correlation"]["panels"][0]["rows"][0]
    outliers = diagnostics["outliers"]
    elasticity = diagnostics["elasticities"]["items"][0]
    coefficient = next(row for row in view["coef_rows"] if row["pvalue"] != "—")
    bound_values = [
        (metrics["r2"], metrics["tokens"]["r2"]),
        (metrics["f_statistic"], metrics["tokens"]["f_statistic"]),
        (metrics["f_pvalue"], metrics["tokens"]["f_pvalue"]),
        (coefficient["pvalue"], coefficient["pvalue_token"]),
        (normal["observed_count"], normal["tokens"]["observed_count"]),
        (matrix["values"][0], matrix["cell_tokens"][0]),
        (outliers["detected"], outliers["tokens"]["detected"]),
        (elasticity["elasticity"], elasticity["tokens"]["elasticity"]),
    ]
    for value, token in bound_values:
        pattern = re.escape(str(value)) + r"\s*" + re.escape(token)
        mutation, count = re.subn(pattern, "999999" + token, text, count=1)
        assert count == 1, (value, token)
        check = verify_report_consistency(
            pdf, snapshot, context, extracted_text=mutation
        )
        assert check["ok"] is False
        assert any(item["code"] == "MUTATED_STATISTICAL_FIELD"
                   for item in check["findings"])


def test_xlsx_sanitizes_all_text_and_rejects_cell_level_tampering():
    snapshot, context = _rich_case()
    first = context["used_rows"][0]
    first["source"] = "=HYPERLINK(\"https://invalid.test\")"
    first["address"] = "+CMD"
    first["justification"] = "@malicioso"
    first["values"]["=cabecalho"] = "=1+1"
    snapshot["job_id"] = "=job"
    data = build_sample_xlsx(snapshot, context)
    assert verify_sample_xlsx(data, snapshot, context)["ok"] is True
    workbook = load_workbook(io.BytesIO(data), data_only=False)
    assert all(
        cell.data_type != "f"
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
    )

    def mutated(change):
        candidate = load_workbook(io.BytesIO(data), data_only=False)
        change(candidate)
        output = io.BytesIO()
        candidate.save(output)
        check = verify_sample_xlsx(output.getvalue(), snapshot, context)
        assert check["ok"] is False
        return {item["code"] for item in check["findings"]}

    assert "SAMPLE_XLSX_CONTENT_MISMATCH" in mutated(
        lambda book: setattr(book["amostra_efetiva"].cell(2, 8), "value",
                             book["amostra_efetiva"].cell(3, 8).value)
    )
    assert "SAMPLE_XLSX_GEOLOCATION_INVALID" in mutated(
        lambda book: setattr(book["amostra_efetiva"].cell(2, 4), "value", 91)
    )
    assert "SAMPLE_XLSX_CONTENT_MISMATCH" in mutated(
        lambda book: setattr(book["amostra_efetiva"].cell(2, 6), "value",
                             "fonte adulterada")
    )
    assert "SAMPLE_XLSX_EXCLUDED_CONTENT_MISMATCH" in mutated(
        lambda book: setattr(book["dados_excluidos"].cell(2, 6), "value",
                             "fonte excluída adulterada")
    )
    assert "SAMPLE_XLSX_FORMULA_CELL" in mutated(
        lambda book: setattr(book["amostra_efetiva"].cell(1, 8), "value", "=1+1")
    )
    assert "SAMPLE_XLSX_ROW_IDS_MISMATCH" in mutated(
        lambda book: book["amostra_efetiva"].delete_rows(2)
    )
    assert "SAMPLE_XLSX_ROW_IDS_MISMATCH" in mutated(
        lambda book: book["amostra_efetiva"].append(
            [cell.value for cell in book["amostra_efetiva"][2]]
        )
    )
    assert "SAMPLE_XLSX_CONTENT_MISMATCH" in mutated(
        lambda book: setattr(book["amostra_efetiva"].cell(2, 1), "value", True)
    )
    assert "SAMPLE_XLSX_SHEET_HIDDEN" in mutated(
        lambda book: setattr(book["amostra_efetiva"], "sheet_state", "hidden")
    )
    assert "SAMPLE_XLSX_ROW_HIDDEN" in mutated(
        lambda book: setattr(book["amostra_efetiva"].row_dimensions[2], "hidden", True)
    )
    assert "SAMPLE_XLSX_COLUMN_HIDDEN" in mutated(
        lambda book: setattr(book["amostra_efetiva"].column_dimensions["B"], "hidden", True)
    )
    assert "SAMPLE_XLSX_CELL_FORMAT_UNEXPECTED" in mutated(
        lambda book: setattr(book["amostra_efetiva"]["H2"], "number_format", ";;;")
    )


def test_test_certificate_configuration_cannot_become_icp_authority():
    fake_anchor = "f" * 64
    assert fake_anchor not in ICP_BRASIL_TRUST_ANCHORS
    snapshot, context = _rich_case()
    xlsx = build_sample_xlsx(snapshot, context)
    context["output_manifest"] = build_output_manifest(snapshot, context)
    pdf = render_report(snapshot, context)
    docx = build_docx(snapshot, context)
    forged_record = {
        "status": "valid", "synthetic_test_only": True,
        "local_verification": {
            "status": "valid",
            "policy": {"trust_framework": "ICP-Brasil",
                       "trust_anchor_allowlist_verified": True},
        },
    }
    manifest = _representation_manifest(
        snapshot, context, pdf=pdf, docx=docx, sample_xlsx=xlsx,
        signed_pdf=pdf + b"\n%TESTE", signature_record=forged_record,
    )
    assert manifest["representations"]["signed_report.pdf"]["icp_brasil_verified"] is False
    assert "bb.guiar.assinatura_icp" not in manifest["items"]

    # A valid-looking generic record bound to the bytes is still insufficient
    # once the frozen recipient profile requires ICP-Brasil.
    signed = pdf + b"\n%TESTE"
    bb_snapshot = copy.deepcopy(snapshot)
    bb_snapshot["provenance"]["qualification_context"]["resolved_profile"]["id"] = (
        "bb-meci-avaliacao-imovel-pf"
    )
    record = {
        "status": "valid", "backend": "pyHanko", "synthetic_test_only": True,
        "unsigned_pdf_sha256": hashlib.sha256(pdf).hexdigest(),
        "signed_pdf_sha256": hashlib.sha256(signed).hexdigest(),
        "snapshot_sha256": signable_snapshot_sha256(bb_snapshot),
        "revision_id": "rev-test-001",
        "result_fingerprint": bb_snapshot["provenance"]["qualification_context"][
            "result_fingerprint"
        ],
        "report_content_fingerprint": report_content_fingerprint(bb_snapshot, context),
        "incremental_base_verified": True,
        "local_verification": {"status": "valid", "signature_count": 1,
                               "policy": {"trust_framework": "ICP-Brasil"}},
    }
    binding = verify_signature_binding(
        signed, record, bb_snapshot, revision_id="rev-test-001",
        unsigned_pdf=pdf, report_context=context,
    )
    assert any(item["code"] == "ICP_BRASIL_SIGNATURE_POLICY_NOT_VERIFIED"
               for item in binding["findings"])


def test_signature_policy_uses_structural_oid_and_authenticated_policy_hash():
    from asn1crypto import cms
    from pyhanko.sign.ades import cades_asn1

    approved_oid = "2.16.76.1.7.1.13.1.4"

    def embedded(main_oid, digest, *, qualifier=None, duplicate=False):
        value = {
            "sig_policy_id": main_oid,
            "sig_policy_hash": {
                "digest_algorithm": {"algorithm": "sha256"},
                "digest": digest,
            },
        }
        if qualifier:
            value["sig_policy_qualifiers"] = [{
                "sig_policy_qualifier_id": "sp_unotice",
                "sig_qualifier": {"explicit_text": ("utf8_string", qualifier)},
            }]
        policy = cades_asn1.SignaturePolicyIdentifier(
            name="signature_policy_id", value=value
        )
        attribute = cms.CMSAttribute({
            "type": "signature_policy_identifier", "values": [policy]
        })

        class SignerInfo:
            def __getitem__(self, key):
                if key == "signed_attrs":
                    return [attribute, attribute] if duplicate else [attribute]
                raise KeyError(key)

        class Embedded:
            signer_info = SignerInfo()

        return Embedded()

    official_digest = bytes.fromhex(ICP_BRASIL_PADES_POLICY_REGISTRY[approved_oid])
    structural = _signature_policy_details(embedded(approved_oid, official_digest))
    assert structural == {
        "oid": approved_oid,
        "hash_algorithm": "sha256",
        "hash_value": official_digest.hex(),
    }
    assert _signature_policy_verified(structural, registry_current=True) is True
    qualifier_injection = _signature_policy_details(
        embedded("1.2.3.4.999", b"X" * 32, qualifier=approved_oid)
    )
    assert qualifier_injection["oid"] == "1.2.3.4.999"
    assert qualifier_injection["hash_value"] != ICP_BRASIL_PADES_POLICY_REGISTRY.get(
        qualifier_injection["oid"]
    )
    assert _signature_policy_verified(qualifier_injection, registry_current=True) is False
    wrong_hash = {**structural, "hash_value": (b"X" * 32).hex()}
    wrong_algorithm = {**structural, "hash_algorithm": "sha512"}
    assert _signature_policy_verified(wrong_hash, registry_current=True) is False
    assert _signature_policy_verified(wrong_algorithm, registry_current=True) is False
    assert _signature_policy_verified(structural, registry_current=False) is False
    assert _signature_policy_details(
        embedded(approved_oid, official_digest, duplicate=True)
    ) is None


def test_resolved_bb_product_baseline_has_capability_without_inventing_case_inputs():
    profile = resolve_profile({"id": "bb-meci-avaliacao-imovel-pf", "version": "0.3.0"})
    baseline = product_conformance_baseline(profile)
    assert baseline["product_can_meet_standard"] is True
    assert baseline["counts"]["partial"] == 0
    assert baseline["counts"]["missing"] == 0
    assert "bb.guiar.assinatura_icp" in baseline["by_state"]["emitted"]
    assert "10.1.c" in baseline["requires_human_input"]
    empty_case = assess_output_conformance(profile, {"items": {}})
    assert empty_case["would_be_accepted_without_reservations"] is False
    assert "bb.guiar.assinatura_icp" in empty_case["awaiting_signature"]


def test_revocation_requires_strict_context_and_accepts_real_test_crl_only_when_current():
    snapshot, context = qualified_case()
    pdf = render_report(snapshot, context)
    signed, strict_context = _signed_with_crl(pdf)
    valid = verify_pdf_signature(
        signed, validation_context=strict_context,
        policy={"require_revocation": True},
    )
    assert valid["status"] == "valid"
    assert valid["policy"]["revocation_policy_verified"] is True
    assert valid["signatures"][0]["revoked"] is False
    request = prepare_signature_request(
        pdf,
        snapshot,
        revision_id="REV-CRL-TESTE",
        profile_id="urban-market-regression",
        report_context=context,
    )
    record = record_external_signature(
        signed,
        request,
        unsigned_pdf=pdf,
        validation_context=strict_context,
        policy={"require_revocation": True},
    )
    assert record["status"] == "valid"
    assert verify_signature_binding(
        signed,
        record,
        snapshot,
        revision_id="REV-CRL-TESTE",
        unsigned_pdf=pdf,
        report_context=context,
    )["ok"] is True
    changed_snapshot = copy.deepcopy(snapshot)
    changed_snapshot["provenance"]["qualification_context"]["rule_results"][0][
        "status"
    ] = "failed"
    assert verify_signature_binding(
        signed,
        record,
        changed_snapshot,
        revision_id="REV-CRL-TESTE",
        unsigned_pdf=pdf,
        report_context=context,
    )["ok"] is False
    changed_context = {**context, "observations": "TESTE: declaração alterada"}
    assert verify_signature_binding(
        signed,
        record,
        snapshot,
        revision_id="REV-CRL-TESTE",
        unsigned_pdf=pdf,
        report_context=changed_context,
    )["ok"] is False
    assert verify_signature_binding(
        signed + b"tamper",
        record,
        snapshot,
        revision_id="REV-CRL-TESTE",
        unsigned_pdf=pdf,
        report_context=context,
    )["ok"] is False

    missing = verify_pdf_signature(
        signed, validation_context=None, policy={"require_revocation": True}
    )
    assert missing["status"] != "valid"
    assert any(item["code"] == "PDF_REVOCATION_POLICY_NOT_STRICT"
               for item in missing["findings"])

    revoked_signed, revoked_context = _signed_with_crl(pdf, revoked=True)
    revoked = verify_pdf_signature(
        revoked_signed, validation_context=revoked_context,
        policy={"require_revocation": True},
    )
    assert revoked["status"] == "invalid"
    assert any(item["code"] == "PDF_SIGNER_CERTIFICATE_REVOKED"
               for item in revoked["findings"])


def test_external_signature_transition_preserves_only_signable_material_identity():
    historical_shapes = [
        {},
        {"result_fingerprint": "historical-result"},
        {"rule_results": [], "result_fingerprint": "historical-result"},
    ]
    for historical in historical_shapes:
        unchanged = copy.deepcopy(historical)
        _normalize_external_signature_state(unchanged)
        assert unchanged == historical

    snapshot, context = _rich_case()
    pending = copy.deepcopy(snapshot)
    qctx = pending["provenance"]["qualification_context"]
    signature_requirement = (
        "Via de assinatura digital do PDF por certificado ICP-Brasil"
    )
    pending_explanation = (
        signature_requirement
        + " — a via de importação e verificação existe, mas os bytes assinados e a "
        "cadeia de confiança ainda não foram fornecidos nesta etapa. A pendência "
        "permite preparar os bytes para assinatura, mas impede o pacote final assinado."
    )
    signature_rule = {
        "rule_id": "bb.guiar.assinatura_icp",
        "source_id": "bb-meci-avaliacao-imovel-pf",
        "edition_or_version": "0.3.0",
        "clause": "GUIAR / assinatura",
        "applicability": "applicable",
        "status": "pending_manual",
        "observed": None,
        "criterion_ref": "external_signature",
        "evidence_refs": [],
        "explanation": pending_explanation,
    }
    signature_evidence = (
        "MP-OUTPUT-MANIFEST/1:representations.signed_report.pdf"
        "#assinatura-cadeia-validada"
    )
    qctx.setdefault("rule_results", []).append(copy.deepcopy(signature_rule))
    qctx["output_conformance"] = {
        "applicable_requirements": 40,
        "conforming": 39,
        "rule_results": [copy.deepcopy(signature_rule)],
        "blocking": [{
            "code": "output_requirement_pending_signature",
            "requirement_id": "bb.guiar.assinatura_icp",
            "clause": "GUIAR / assinatura",
            "detail": pending_explanation,
            "owner": "profissional responsável / cadeia de confiança configurada",
            "stage": "post_review_external_signature",
        }],
        "awaiting_signature": ["bb.guiar.assinatura_icp"],
        "would_be_accepted_without_reservations": False,
    }

    signed = copy.deepcopy(pending)
    signed_qctx = signed["provenance"]["qualification_context"]
    for rule in signed_qctx["rule_results"]:
        if rule.get("rule_id") == "bb.guiar.assinatura_icp":
            rule.update({
                "status": "passed",
                "observed": signature_evidence,
                "evidence_refs": [signature_evidence],
                "explanation": signature_requirement,
            })
    output = signed_qctx["output_conformance"]
    output["rule_results"][0].update({
        "status": "passed",
        "observed": signature_evidence,
        "evidence_refs": [signature_evidence],
        "explanation": signature_requirement,
    })
    output.update({
        "conforming": 40,
        "blocking": [],
        "awaiting_signature": [],
        "would_be_accepted_without_reservations": True,
    })
    signed_qctx["digital_signature"] = {"status": "valid"}
    signed_qctx["signed_bytes_sha256"] = "a" * 64
    assert report_content_fingerprint(pending, context) == report_content_fingerprint(
        signed, context
    )
    assert signable_snapshot_sha256(pending) == signable_snapshot_sha256(signed)

    changed_rule = copy.deepcopy(signed)
    for rule in changed_rule["provenance"]["qualification_context"]["rule_results"]:
        if rule.get("rule_id") != "bb.guiar.assinatura_icp":
            rule["status"] = "failed"
            break
    assert report_content_fingerprint(changed_rule, context) != report_content_fingerprint(
        signed, context
    )
    assert signable_snapshot_sha256(changed_rule) != signable_snapshot_sha256(signed)
    changed_declaration = copy.deepcopy(context)
    changed_declaration["observations"] += " alteração material"
    assert report_content_fingerprint(signed, changed_declaration) != report_content_fingerprint(
        signed, context
    )

    invalid_signature_state = copy.deepcopy(pending)
    invalid_qctx = invalid_signature_state["provenance"]["qualification_context"]
    for rule in invalid_qctx["rule_results"]:
        if rule.get("rule_id") == "bb.guiar.assinatura_icp":
            rule["status"] = "failed"
    invalid_output = invalid_qctx["output_conformance"]
    invalid_output["rule_results"][0]["status"] = "failed"
    invalid_output["blocking"][0]["code"] = "output_requirement_partial"
    assert report_content_fingerprint(
        invalid_signature_state, context
    ) != report_content_fingerprint(pending, context)
    assert signable_snapshot_sha256(invalid_signature_state) != signable_snapshot_sha256(
        pending
    )

    additional_stage = copy.deepcopy(pending)
    additional_stage["provenance"]["qualification_context"]["output_conformance"][
        "awaiting_signature"
    ].append("other.external.signature")
    assert report_content_fingerprint(additional_stage, context) != report_content_fingerprint(
        pending, context
    )
    assert signable_snapshot_sha256(additional_stage) != signable_snapshot_sha256(pending)

    malformed_evidence = copy.deepcopy(signed)
    malformed_qctx = malformed_evidence["provenance"]["qualification_context"]
    for rule in malformed_qctx["rule_results"]:
        if rule.get("rule_id") == "bb.guiar.assinatura_icp":
            rule["evidence_refs"] = signature_evidence
    malformed_qctx["output_conformance"]["rule_results"][0][
        "evidence_refs"
    ] = signature_evidence
    assert report_content_fingerprint(
        malformed_evidence, context
    ) != report_content_fingerprint(signed, context)
    assert signable_snapshot_sha256(malformed_evidence) != signable_snapshot_sha256(
        signed
    )
    for aggregate_mutation in ("duplicate_blocker", "duplicate_awaiting", "detail"):
        malformed_pending = copy.deepcopy(pending)
        malformed_output = malformed_pending["provenance"]["qualification_context"][
            "output_conformance"
        ]
        if aggregate_mutation == "duplicate_blocker":
            malformed_output["blocking"].append(
                copy.deepcopy(malformed_output["blocking"][0])
            )
        elif aggregate_mutation == "duplicate_awaiting":
            malformed_output["awaiting_signature"].append(
                "bb.guiar.assinatura_icp"
            )
        else:
            malformed_output["blocking"][0]["detail"] += " adulterado"
        assert report_content_fingerprint(
            malformed_pending, context
        ) != report_content_fingerprint(pending, context)
        assert signable_snapshot_sha256(
            malformed_pending
        ) != signable_snapshot_sha256(pending)
    mixed_phase = copy.deepcopy(pending)
    mixed_top = next(
        rule for rule in mixed_phase["provenance"]["qualification_context"][
            "rule_results"
        ] if rule.get("rule_id") == "bb.guiar.assinatura_icp"
    )
    mixed_top.update({
        "status": "passed",
        "observed": signature_evidence,
        "evidence_refs": [signature_evidence],
        "explanation": signature_requirement,
    })
    assert report_content_fingerprint(mixed_phase, context) != report_content_fingerprint(
        pending, context
    )
    assert signable_snapshot_sha256(mixed_phase) != signable_snapshot_sha256(pending)


def test_professional_review_is_invalidated_by_report_content_change_without_signature():
    from tests.comercial.test_c06_qualification_integration import _compose, _spec

    first_report = "a" * 64
    changed_report = "b" * 64
    first = _compose(_spec(report_content_fingerprint=first_report))
    review = {
        "professional_id": "CREA-TESTE-000",
        "motive": "TESTE: revisão do conteúdo A",
        "version": "REV-TESTE-1",
        "fingerprint": first["result_fingerprint"],
        "report_content_fingerprint": first_report,
    }
    reviewed = _compose(_spec(
        report_content_fingerprint=first_report, review_events=[review]
    ))
    changed = _compose(_spec(
        report_content_fingerprint=changed_report, review_events=[review]
    ))
    assert reviewed["case_release_status"] == "ready_for_professional_signoff"
    assert changed["result_fingerprint"] == reviewed["result_fingerprint"]
    assert changed["case_release_status"] == "review_required"
    assert changed["stale_review_events"] == [review]
    review_without_report = {key: value for key, value in review.items()
                             if key != "report_content_fingerprint"}
    missing_binding = _compose(_spec(
        report_content_fingerprint=first_report,
        review_events=[review_without_report],
    ))
    assert missing_binding["case_release_status"] == "review_required"
    assert missing_binding["stale_review_events"] == [review_without_report]
    historical = _compose(_spec(review_events=[review_without_report]))
    assert historical["case_release_status"] == "ready_for_professional_signoff"


def test_real_worker_numeric_disclosure_reaches_verified_pdf_docx_xlsx_bytes(tmp_path):
    csv_bytes, _areas, _prices = _noisy_linear_csv()
    profile = resolve_profile({"id": "bb-meci-avaliacao-imovel-pf", "version": "0.3.0"})
    spec = _professional_spec()
    spec["qualification_profile"] = {
        key: profile[key]
        for key in (
            "id", "version", "source_set_sha256", "purpose", "value_basis",
            "method", "asset_scope", "recipient_id",
        )
    }
    spec["report_context"] = {
        "objective": "TESTE: estimar valor de mercado.",
        "market_diagnosis": "TESTE: mercado sintético para integração.",
        "variable_classification": {
            "area": {"criterion": "área em m²", "coding": "contínua"}
        },
        "observations": "TESTE: intervalos calculados registrados.",
        "professional_identity": {
            "name": "Profissional TESTE", "council": "CREA-TESTE",
            "registration": "000000", "art_rrt": "ART-TESTE-000",
        },
        "sample_evidence": {
            f"R{index - 1:06d}": {
                "address": f"Rua Integração, {index}, Cidade TESTE/TS",
                "latitude": -27.5 - index / 10000,
                "longitude": -48.5 - index / 10000,
                "source": f"TESTE: fonte {index}",
            }
            for index in range(1, 31)
        },
    }
    subject = {
        "area": 73.5, "bairro": "Centro",
        "geolocation": {
            "address": "Rua Avaliando, 1, Cidade TESTE/TS",
            "latitude": -27.6, "longitude": -48.6,
            "source": "TESTE: declaração de vistoria",
        },
    }
    store = JobStore(tmp_path / "store", recover_abandoned=False)
    created = store.create(request_spec=spec, payload={"filename": "SINTETICO.csv"})
    product = compose_valuation_job(
        job_id=created["job_id"], file_bytes=csv_bytes, filename="SINTETICO.csv",
        request_spec=spec, subject_raw=subject, project_id=None,
        peers=resolve_peers(), job_store=store,
    )
    snapshot = product["snapshot"]
    context = __import__("json").loads(
        store.get_artifact(created["job_id"], "report_context.json")
    )
    diagnostics = snapshot["validation"]["statistical"]["diagnostics"]
    assert diagnostics["standardized_residuals"]["available"] is True
    assert diagnostics["correlation_matrix"]["status"] == "complete"
    assert diagnostics["elasticities"]["coverage"]["complete"] is True
    assert diagnostics["outlier_count"]["coverage"]["complete"] is True

    xlsx = build_sample_xlsx(snapshot, context)
    xlsx_check = verify_sample_xlsx(xlsx, snapshot, context)
    assert xlsx_check["ok"] is True
    assert xlsx_check["geolocation_complete"] is True
    content_manifest = build_output_manifest(snapshot, context, representations={
        "sample.xlsx": {"verified": True, "sha256": xlsx_check["sha256"],
                        "size": xlsx_check["size"]}
    })
    context["output_manifest"] = content_manifest
    pdf = render_report(snapshot, context)
    docx = build_docx(snapshot, context)
    representation = _representation_manifest(
        snapshot, context, pdf=pdf, docx=docx, sample_xlsx=xlsx
    )
    check = _verify_representation_manifest(
        representation, snapshot, context,
        {"report.pdf": pdf, "report.docx": docx, "sample.xlsx": xlsx},
    )
    assert check["ok"] is True, check
    required_worker_items = {
        "meci.3.3.1.b", "meci.3.3.1.c", "meci.3.3.1.j", "meci.3.3.1.k",
        "meci.3.3.1.l", "meci.3.3.1.m1", "meci.3.3.1.m2",
        "meci.2.3.3.2.excel", "normas.12.1.8.geolocalizacao",
    }
    assert required_worker_items <= set(representation["items"])
    text = extract_pdf_text(pdf)
    assert "Coeficiente de correlação R" in text
    assert "Matriz de correlações" in text
    assert "Dados discrepantes e influentes:" in text
    assert "resíduo internamente studentizado" in text
    inspection = Path("/tmp/c06-output-conformance-inspection")
    inspection.mkdir(parents=True, exist_ok=True)
    artifacts = {
        "report.pdf": pdf,
        "report.docx": docx,
        "sample.xlsx": xlsx,
        "snapshot.json": json.dumps(
            snapshot, ensure_ascii=False, sort_keys=True, indent=2
        ).encode("utf-8"),
        "report_context.json": json.dumps(
            context, ensure_ascii=False, sort_keys=True, indent=2
        ).encode("utf-8"),
        "output_representation_manifest.json": json.dumps(
            representation, ensure_ascii=False, sort_keys=True, indent=2
        ).encode("utf-8"),
    }
    inventory = {
        name: {"sha256": hashlib.sha256(data).hexdigest(), "size": len(data)}
        for name, data in artifacts.items()
    }
    for name, data in artifacts.items():
        (inspection / name).write_bytes(data)
    (inspection / "inventory.json").write_text(
        json.dumps({"artifacts": inventory, "manifest_verification": check},
                   ensure_ascii=False, sort_keys=True, indent=2),
        encoding="utf-8",
    )
